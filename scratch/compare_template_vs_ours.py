import docx

doc_tpl = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')
doc_our = docx.Document('assets/jurnal/JURNAL_SIMASI_TRUSTGIVE.docx')

print("=== TEMPLATE TITLE & HEADER ===")
for p in doc_tpl.paragraphs[:16]:
    if p.text.strip():
        fonts = [(r.text[:20], r.font.name, r.font.size.pt if r.font.size else None, r.bold, r.italic) for r in p.runs]
        pPr = p._p.pPr
        jc = pPr.xpath('./w:jc/@w:val') if pPr is not None else []
        ind = pPr.xpath('./w:ind') if pPr is not None else []
        ind_d = ind[0].attrib if ind else {}
        sp = pPr.xpath('./w:spacing') if pPr is not None else []
        sp_d = sp[0].attrib if sp else {}
        print(f"[{p.style.name}] (jc={jc}, ind={ind_d}, sp={sp_d}): {p.text[:60]}")
        print(f"   Runs: {fonts[:3]}")

print("\n=== OUR TITLE & HEADER ===")
for p in doc_our.paragraphs[:16]:
    if p.text.strip():
        fonts = [(r.text[:20], r.font.name, r.font.size.pt if r.font.size else None, r.bold, r.italic) for r in p.runs]
        pPr = p._p.pPr
        jc = pPr.xpath('./w:jc/@w:val') if pPr is not None else []
        ind = pPr.xpath('./w:ind') if pPr is not None else []
        ind_d = ind[0].attrib if ind else {}
        sp = pPr.xpath('./w:spacing') if pPr is not None else []
        sp_d = sp[0].attrib if sp else {}
        print(f"[{p.style.name}] (jc={jc}, ind={ind_d}, sp={sp_d}): {p.text[:60]}")
        print(f"   Runs: {fonts[:3]}")

print("\n=== TEMPLATE TABLES DETAILS ===")
for i, t in enumerate(doc_tpl.tables[1:3]):
    print(f"Table {i+1} style: {t.style.name if t.style else None}, alignment: {t.alignment}")
    tblPr = t._tbl.tblPr
    tblBorders = tblPr.xpath('./w:tblBorders') if tblPr is not None else []
    if tblBorders:
        for b in tblBorders[0]:
            print(f"  border {b.tag.split('}')[-1]}: {b.attrib}")
    # Inspect first cell
    cell = t.cell(0, 0)
    tcPr = cell._tc.tcPr
    shd = tcPr.xpath('./w:shd/@w:fill') if tcPr is not None else []
    tcBorders = tcPr.xpath('./w:tcBorders') if tcPr is not None else []
    print(f"  Cell 0,0: text='{cell.text.strip()}', shd={shd}, borders={len(tcBorders)}")
    p = cell.paragraphs[0]
    r_info = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
    print(f"  Cell 0,0 text runs: {r_info}")

print("\n=== OUR TABLES DETAILS ===")
for i, t in enumerate(doc_our.tables[1:3]):
    print(f"Table {i+1} style: {t.style.name if t.style else None}, alignment: {t.alignment}")
    cell = t.cell(0, 0)
    tcPr = cell._tc.tcPr
    shd = tcPr.xpath('./w:shd/@w:fill') if tcPr is not None else []
    print(f"  Cell 0,0: text='{cell.text.strip()}', shd={shd}")
    p = cell.paragraphs[0]
    r_info = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
    print(f"  Cell 0,0 text runs: {r_info}")
