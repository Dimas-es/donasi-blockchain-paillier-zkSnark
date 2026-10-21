import docx

doc = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')

print("=== TOTAL PARAGRAPHS:", len(doc.paragraphs))
print("=== TOTAL TABLES:", len(doc.tables))

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    pPr = p._p.pPr
    jc = pPr.xpath('./w:jc/@w:val') if pPr is not None else []
    ind = pPr.xpath('./w:ind') if pPr is not None else []
    ind_attr = ind[0].attrib if ind else {}
    sp = pPr.xpath('./w:spacing') if pPr is not None else []
    sp_attr = sp[0].attrib if sp else {}
    
    runs_details = []
    for r in p.runs:
        if r.text:
            runs_details.append({
                'text': r.text[:30],
                'font': r.font.name,
                'size': r.font.size.pt if r.font.size else None,
                'bold': r.bold,
                'italic': r.italic,
                'underline': r.underline
            })
    print(f"\n--- P{i:03d} [{p.style.name}] align={jc} ind={ind_attr} sp={sp_attr} ---")
    print(f"Text ({len(txt)} chars): {txt[:120]}...")
    print(f"Runs: {runs_details}")
