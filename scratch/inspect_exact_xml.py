import docx

doc = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')

def print_p_xml(idx, label):
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        print(f"\n=== {label} (P{idx:02d}) ===")
        print(f"Style: {p.style.name}, Text: {p.text[:60]}")
        pPr = p._p.pPr
        if pPr is not None:
            print("pPr XML:", pPr.xml)
        for r_idx, r in enumerate(p.runs[:2]):
            rPr = r._r.rPr
            if rPr is not None:
                print(f"  run {r_idx} rPr XML: {rPr.xml}")

print_p_xml(0, "TITLE 1")
print_p_xml(3, "AUTHORS")
print_p_xml(5, "AFFILIATION")
print_p_xml(6, "EMAIL")
print_p_xml(8, "ABSTRACT HEADER")
print_p_xml(9, "ABSTRACT BODY")
print_p_xml(10, "KEYWORDS")
print_p_xml(12, "ABSTRAK HEADER")
print_p_xml(13, "ABSTRAK BODY")
print_p_xml(14, "KATA KUNCI")
print_p_xml(16, "HEADING 1 (PENDAHULUAN)")
print_p_xml(17, "BODY TEXT")
print_p_xml(21, "CAPTION GAMBAR")
print_p_xml(72, "CAPTION TABEL")
