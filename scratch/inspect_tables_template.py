import docx

doc = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')

print(f"Total tables: {len(doc.tables)}")

for t_idx in range(len(doc.tables)):
    t = doc.tables[t_idx]
    print(f"\n=================== TABLE {t_idx} ===================")
    print(f"Rows: {len(t.rows)}, Cols: {len(t.columns)}, Alignment: {t.alignment}")
    
    # Check borders from XML
    tblPr = t._tbl.tblPr
    tblBorders = tblPr.xpath('./w:tblBorders') if tblPr is not None else []
    if tblBorders:
        for b in tblBorders[0]:
            print(f"  tblBorder {b.tag.split('}')[-1]}: {b.attrib}")
    else:
        print("  no w:tblBorders on table")
        
    # Check header row (row 0)
    for c_idx, cell in enumerate(t.rows[0].cells[:4]):
        p = cell.paragraphs[0]
        p_text = p.text.strip()
        p_align = p.alignment
        runs_info = [(r.text[:15], r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
        shd = cell._tc.xpath('./w:tcPr/w:shd/@w:fill')
        borders = cell._tc.xpath('./w:tcPr/w:tcBorders')
        print(f"  Col {c_idx}: text='{p_text}', align={p_align}, shd={shd}, borders={len(borders)}, runs={runs_info}")
        
    if len(t.rows) > 1:
        # Check row 1
        for c_idx, cell in enumerate(t.rows[1].cells[:4]):
            p = cell.paragraphs[0]
            p_text = p.text.strip()
            runs_info = [(r.text[:15], r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
            shd = cell._tc.xpath('./w:tcPr/w:shd/@w:fill')
            print(f"  Row 1 Col {c_idx}: text='{p_text}', shd={shd}, runs={runs_info}")
