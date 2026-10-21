import docx
from docx.oxml import xmlchemy

doc = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')
print('=== SECTIONS ===')
for i, s in enumerate(doc.sections):
    print(f'Section {i}: start_type={s.start_type}, top={s.top_margin.cm:.2f}cm, bot={s.bottom_margin.cm:.2f}cm, left={s.left_margin.cm:.2f}cm, right={s.right_margin.cm:.2f}cm, page_w={s.page_width.cm:.2f}cm, page_h={s.page_height.cm:.2f}cm')
    sectPr = s._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        print(f'  cols attrib: {cols[0].attrib}')
    else:
        print('  no w:cols found')

print('\n=== PARAGRAPHS IN TEMPLATE ===')
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        runs_info = []
        for r in p.runs:
            sz = r.font.size.pt if r.font.size else None
            fn = r.font.name
            b = r.bold
            it = r.italic
            runs_info.append((r.text[:30], fn, sz, b, it))
        pPr = p._p.pPr
        jc = pPr.xpath('./w:jc/@w:val') if pPr is not None else []
        ind = pPr.xpath('./w:ind') if pPr is not None else []
        ind_str = ind[0].attrib if ind else {}
        sp = pPr.xpath('./w:spacing') if pPr is not None else []
        sp_str = sp[0].attrib if sp else {}
        print(f'P{i:03d} [{p.style.name}] jc={jc} ind={ind_str} sp={sp_str}')
        print(f'     Text: {p.text[:80]}')
        print(f'     Runs: {runs_info[:3]}')

print('\n=== TABLES IN TEMPLATE ===')
for i, t in enumerate(doc.tables):
    print(f'Table {i}: rows={len(t.rows)}, cols={len(t.columns)}')
    for r_idx, row in enumerate(t.rows[:3]):
        row_txt = [c.text.strip()[:30] for c in row.cells]
        print(f'  Row {r_idx}: {row_txt}')
