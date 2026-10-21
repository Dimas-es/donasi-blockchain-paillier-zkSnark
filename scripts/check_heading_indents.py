import docx

doc_tpl = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')
doc_gen = docx.Document('assets/jurnal/JURNAL_SIMASI_TRUSTGIVE.docx')

print('=== TEMPLATE P22 (METODE PENELITIAN) & P23 (Pengumpulan Data) ===')
for i in [22, 23, 24, 31, 32]:
    p = doc_tpl.paragraphs[i]
    pPr = p._p.pPr
    ind = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ind') if pPr is not None else None
    print(f'TPL P{i} [{p.style.name}]: text="{p.text[:40]}"')
    if ind is not None:
        print('   ind:', ind.attrib)
    else:
        print('   ind: None')

print('\n=== GENERATED METODE PENELITIAN & SUBHEADING ===')
for i, p in enumerate(doc_gen.paragraphs):
    if 'METODE PENELITIAN' in p.text or 'Arsitektur Hibrida' in p.text or 'Penelitian ini dilaksanakan' in p.text:
        pPr = p._p.pPr
        ind = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ind') if pPr is not None else None
        print(f'GEN P{i} [{p.style.name}]: text="{p.text[:40]}"')
        if ind is not None:
            print('   ind:', ind.attrib)
        else:
            print('   ind: None')
