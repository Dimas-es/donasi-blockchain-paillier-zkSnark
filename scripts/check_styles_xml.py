import docx

doc = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')

print('=== STYLES IN TEMPLATE ===')
for s in doc.styles:
    if s.name in ['Title', 'Heading 1', 'Heading 2', 'Body Text', 'Caption', 'Normal', 'List Paragraph']:
        font = s.font
        color = font.color.rgb if font.color else 'None'
        print(f'Style: {s.name:15} | font={font.name}, size={font.size.pt if font.size else None}, bold={font.bold}, italic={font.italic}, color={color}')

# Check paragraph 3 in template (Author line)
p3 = doc.paragraphs[3]
print('\nP3 Author line XML:')
print(p3._p.xml)

# Check a Caption in template
for i, p in enumerate(doc.paragraphs):
    if p.text.startswith('Gambar') or p.text.startswith('Tabel'):
        print(f'\nP{i} Caption XML ({p.text[:30]}):')
        print(p._p.xml)
        break
