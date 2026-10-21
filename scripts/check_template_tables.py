import docx

doc = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')
print('Template tables count:', len(doc.tables))
for i, tbl in enumerate(doc.tables):
    rows = len(tbl.rows)
    cols = len(tbl.columns)
    first_cell = tbl.rows[0].cells[0].text.replace('\n', ' ')[:40]
    print(f'Table {i}: {rows}x{cols}, first cell="{first_cell}"')
