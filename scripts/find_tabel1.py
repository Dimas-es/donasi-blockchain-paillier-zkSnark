import docx

doc = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')
for i, p in enumerate(doc.paragraphs):
    if 'Tabel 1' in p.text:
        print(f'P{i}: "{p.text}"')
        print(f'Previous P{i-1}: "{doc.paragraphs[i-1].text[:80]}"')
        print(f'Next P{i+1}: "{doc.paragraphs[i+1].text[:80]}"')
