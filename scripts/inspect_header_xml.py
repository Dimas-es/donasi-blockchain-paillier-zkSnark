import docx

doc = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')
sec = doc.sections[0]

print("=== TEMPLATE SECTION 0 ===")
print("Different first page:", sec.different_first_page_header_footer)
print("Header paragraphs:")
for p in sec.header.paragraphs:
    print("  header:", repr(p.text))
print("Footer paragraphs:")
for p in sec.footer.paragraphs:
    print("  footer:", repr(p.text))

# Check XML of header
header_xml = sec.header._element.xml
print("Header XML length:", len(header_xml))
