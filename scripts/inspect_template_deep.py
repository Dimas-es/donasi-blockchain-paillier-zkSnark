import docx

doc = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')

print('=== TOTAL PARAGRAPHS IN TEMPLATE ===', len(doc.paragraphs))
for i in range(min(40, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f'P{i} [{p.style.name}]: text="{p.text[:70]}"')
        print(f'   align={p.alignment}, left={p.paragraph_format.left_indent}, right={p.paragraph_format.right_indent}, first_line={p.paragraph_format.first_line_indent}, ls={p.paragraph_format.line_spacing}')
        for r in p.runs[:4]:
            color = r.font.color.rgb if r.font.color else 'None'
            print(f'      run: text="{r.text[:30]}", font={r.font.name}, sz={r.font.size.pt if r.font.size else None}, bold={r.bold}, italic={r.italic}, sup={r.font.superscript}, color={color}')

print('\n=== SECTIONS & HEADERS ===')
for s_idx, s in enumerate(doc.sections):
    print(f'Section {s_idx}: different_first_page={s.different_first_page_header_footer}')
    print('First page header paragraphs:', [p.text for p in s.first_page_header.paragraphs if p.text.strip()])
    print('Header paragraphs:', [p.text for p in s.header.paragraphs if p.text.strip()])
    print('Footer paragraphs:', [p.text for p in s.footer.paragraphs if p.text.strip()])
    print('First page footer paragraphs:', [p.text for p in s.first_page_footer.paragraphs if p.text.strip()])
