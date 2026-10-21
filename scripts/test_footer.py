import docx

doc = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')
sec = doc.sections[0]
print("Footer paragraphs:", [p.text for p in sec.footer.paragraphs])
for i, p in enumerate(sec.footer.paragraphs):
    print(f"Footer P{i} xml:", p._p.xml)
