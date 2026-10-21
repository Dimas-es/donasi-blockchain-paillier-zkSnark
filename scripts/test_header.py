import docx
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document('assets/template/TEMPLATE SIMASI CONTOH.docx')

# Let's inspect the exact XML of Table 0 in doc
t0 = doc.tables[0]
print("Table 0 XML length:", len(t0._tbl.xml))

# Let's inspect running header text in doc
for s in doc.sections:
    print("Header text:", [p.text for p in s.header.paragraphs if p.text.strip()])
