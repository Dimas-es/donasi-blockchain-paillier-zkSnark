import docx

doc = docx.Document('assets/jurnal/JURNAL_SIMASI_TRUSTGIVE.docx')

print("=== CHECKING SECTIONS & HEADERS ===")
for i, s in enumerate(doc.sections):
    hdr_texts = [p.text.strip() for p in s.header.paragraphs if p.text.strip()]
    ftr_texts = [p.text.strip() for p in s.footer.paragraphs if p.text.strip()]
    print(f"Section {i}: diff_first={s.different_first_page_header_footer}")
    print(f"   Header paragraphs: {hdr_texts}")
    print(f"   Footer paragraphs: {ftr_texts}")

print("\n=== CHECKING FIRST 15 PARAGRAPHS (PAGE 1) ===")
for i in range(min(15, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    if p.text.strip():
        runs_info = [(r.text[:25], r.font.name, r.font.size.pt if r.font.size else None, r.bold, r.italic, r.font.superscript, r.font.color.rgb if r.font.color else None) for r in p.runs[:3]]
        print(f"P{i:2d} [{p.style.name:12}] (align={p.alignment}): text='{p.text[:50]}' | runs={runs_info}")

print("\n=== CHECKING ALL CAPTIONS (FIGURES & TABLES) ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.startswith("Gambar ") or p.text.startswith("Tabel "):
        runs_info = [(r.text[:20], r.font.size.pt if r.font.size else None, r.bold, r.italic, r.font.color.rgb if r.font.color else None) for r in p.runs[:2]]
        print(f"Caption P{i:2d}: '{p.text[:55]}' | runs={runs_info}")

print("\n=== CHECKING EMBEDDED IMAGES ===")
img_parts = [p for p in doc.part.related_parts.values() if 'image' in p.content_type]
print(f"Total image parts embedded: {len(img_parts)}")
for idx, img in enumerate(img_parts):
    print(f"  Img {idx}: {img.partname}, size: {len(img.blob)} bytes")
